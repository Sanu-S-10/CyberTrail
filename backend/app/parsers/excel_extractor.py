"""Dynamic Excel Extractor for NCRP Workbooks."""

from __future__ import annotations

import io
import re
from typing import Any
import openpyxl


def _clean_str(val: Any) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    return re.sub(r"\s+", " ", re.sub(r"[\uE000-\uF8FF]", "", s)).strip()


def _parse_amount(val: Any) -> float:
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).replace(",", "").strip()
    match = re.search(r"[-+]?\d*\.\d+|\d+", s)
    return float(match.group()) if match else 0.0


def _parse_layer(val: Any) -> int | None:
    if val is None:
        return None
    if isinstance(val, int):
        return val
    s = str(val).strip()
    if s.isdigit():
        return int(s)
    match = re.search(r"\blayer\s*[:\-]?\s*(\d+)\b", s, re.IGNORECASE)
    if match:
        return int(match.group(1))
    return None


class ExcelExtractor:
    """Extracts NCRP workbook data into three independent collections.

    - ``transfers``: account-to-account relationships. Only sheets that
      explicitly describe a transfer (``Money Transfer to`` and equivalents)
      can produce Money Trail edges; other sheets are never promoted to edges.
    - ``withdrawals``: ATM / POS / Cheque / AEPS cash-out records.
    - ``other_records``: rows from sheets such as *Transaction put on hold*,
      *Other* and *Others Less Then 500*. They are preserved as evidence but
      never become Money Trail transfer relationships.
    """

    # Sheet-name hints that explicitly represent account-to-account transfers.
    TRANSFER_SHEET_MARKERS = ("money transfer", "account transfer", "transfer to", "transaction transfer")
    NON_TRANSFER_SHEET_MARKERS = (
        "put on hold",
        "less then 500",
        "less than 500",
        "withdrawal",
        "atm",
        "pos",
        "cheque",
        "aeps",
        "cash",
    )

    def is_transfer_sheet(self, sheet_name: str) -> bool:
        """True only when the sheet explicitly describes money transfers."""
        name = sheet_name.strip().lower()
        if any(marker in name for marker in self.NON_TRANSFER_SHEET_MARKERS):
            return False
        return any(marker in name for marker in self.TRANSFER_SHEET_MARKERS)

    def process_excel(self, excel_bytes: bytes) -> dict[str, Any]:
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)

        all_transfers: list[dict[str, Any]] = []
        all_withdrawals: list[dict[str, Any]] = []
        other_records: list[dict[str, Any]] = []
        sections: list[dict[str, Any]] = []
        sheets_meta: list[str] = wb.sheetnames

        # Dynamically extract the complainant-reported total fraud amount from any cell.
        reported_fraud_amount = 0.0
        for _sn in wb.sheetnames:
            for _row in wb[_sn].iter_rows(values_only=True):
                _text = " ".join([_clean_str(c) for c in _row if c is not None])
                if not _text:
                    continue
                _m = re.search(r"(?:total\s+)?fraud(?:ulent)?\s*amount[^0-9]{0,60}?([0-9][0-9,]*(?:\.\d+)?)", _text, re.IGNORECASE)
                if _m:
                    reported_fraud_amount = _parse_amount(_m.group(1))
                    break
            if reported_fraud_amount:
                break

        for sheet_idx, sheet_name in enumerate(wb.sheetnames, start=1):
            ws = wb[sheet_name]
            sheet_rows = list(ws.iter_rows(values_only=True))
            if not sheet_rows:
                continue

            # Locate header row
            header_idx = -1
            headers: list[str] = []
            for r_idx, row in enumerate(sheet_rows[:25]):
                row_str = " ".join([_clean_str(c).lower() for c in row if c is not None])
                if any(k in row_str for k in ("account", "amount", "withdrawal", "utr", "s no", "s.no", "transaction")):
                    header_idx = r_idx
                    headers = [_clean_str(c).lower() for c in row]
                    break

            if header_idx == -1:
                continue

            data_rows = sheet_rows[header_idx + 1:]
            sections.append({
                "section_name": sheet_name,
                "sheet_name": sheet_name,
                "row_count": len(data_rows),
                "page_number": sheet_idx,
            })

            # Classify sheet type
            s_name_lower = sheet_name.lower()
            is_withdrawal_sheet = any(w in s_name_lower for w in ("withdrawal", "atm", "pos", "cheque", "aeps", "cash"))
            is_transfer_sheet = self.is_transfer_sheet(sheet_name)

            # Determine column positions
            src_acc_col = self._find_col(headers, ["account no./ (wallet /pg/pa) id", "from account", "source account", "victim account", "debited account"])
            dst_acc_col = self._find_col(headers, ["account no", "to account", "destination account", "credited account", "beneficiary account"])
            
            # If src and dst col point to same index, refine dst_acc_col
            if src_acc_col != -1 and dst_acc_col == src_acc_col:
                # Look for second account no column
                dst_acc_col = self._find_col(headers[src_acc_col + 1:], ["account no", "to account"])
                if dst_acc_col != -1:
                    dst_acc_col += src_acc_col + 1

            bank_col = self._find_col(headers, ["bank/fis", "action taken by bank", "bank name", "bank"])
            layer_col = self._find_col(headers, ["layer", "layers"])
            ifsc_col = self._find_col(headers, ["ifsc code", "ifsc"])
            utr_col = self._find_col(headers, ["transaction id / utr number", "transaction id / utr number2", "reference no", "utr", "approval code"])
            amt_col = self._find_col(headers, ["transaction amount", "withdrawal amount", "disputed amount", "amount"])
            disputed_col = self._find_col(headers, ["disputed amount"])
            date_col = self._find_col(headers, ["transaction date", "withdrawal date & time", "withdrawal date", "date of action", "date"])
            remarks_col = self._find_col(headers, ["remarks", "place/location of atm", "branch location", "merchant name"])

            for r in data_rows:
                if not any(r):
                    continue

                # Check if this row is valid data (e.g. S.No or valid account)
                row_cells = [_clean_str(c) for c in r]
                first_cell = row_cells[0] if row_cells else ""

                # Skip repeated in-table header echoes / blank header-like rows
                first_stripped = first_cell.lower().rstrip('.')
                if first_stripped in ("s no", "s.no", "sno", "s. no"):
                    continue

                # Extract fields
                src_acc = row_cells[src_acc_col] if src_acc_col != -1 and src_acc_col < len(row_cells) else ""
                dst_acc = row_cells[dst_acc_col] if dst_acc_col != -1 and dst_acc_col < len(row_cells) else ""

                # Positional account columns only; a numeric scan fallback would
                # invent endpoints for sheets that are not transfer sheets.
                if not src_acc and not dst_acc and is_transfer_sheet:
                    account_candidates = [c for c in row_cells if len(c) >= 6 and c.isdigit()]
                    if len(account_candidates) >= 2:
                        src_acc, dst_acc = account_candidates[0], account_candidates[1]
                    elif len(account_candidates) == 1:
                        src_acc = account_candidates[0]

                if not src_acc and not dst_acc and not is_withdrawal_sheet and not is_transfer_sheet:
                    # Preserve header-less / summary rows of other sheets as evidence
                    text_cells = [c for c in row_cells if c]
                    if text_cells:
                        other_records.append({
                            "record_type": sheet_name,
                            "sheet_name": sheet_name,
                            "row_text": " | ".join(text_cells),
                        })
                    continue

                if not src_acc and not dst_acc:
                    continue

                amt = _parse_amount(row_cells[amt_col]) if amt_col != -1 and amt_col < len(row_cells) else 0.0
                disputed = _parse_amount(row_cells[disputed_col]) if disputed_col != -1 and disputed_col < len(row_cells) else 0.0
                bank = row_cells[bank_col] if bank_col != -1 and bank_col < len(row_cells) else ""
                layer = _parse_layer(row_cells[layer_col]) if layer_col != -1 and layer_col < len(row_cells) else None
                ifsc = row_cells[ifsc_col] if ifsc_col != -1 and ifsc_col < len(row_cells) else ""
                utr = row_cells[utr_col] if utr_col != -1 and utr_col < len(row_cells) else ""
                date_val = row_cells[date_col] if date_col != -1 and date_col < len(row_cells) else ""
                remarks = row_cells[remarks_col] if remarks_col != -1 and remarks_col < len(row_cells) else ""

                if is_withdrawal_sheet or (src_acc and not dst_acc and ("atm" in s_name_lower or "pos" in s_name_lower or "cheque" in s_name_lower or "aeps" in s_name_lower)):
                    # Determine withdrawal type
                    w_type = "WITHDRAWAL"
                    if "atm" in s_name_lower or "atm" in remarks.lower():
                        w_type = "ATM_WITHDRAWAL"
                    elif "pos" in s_name_lower or "pos" in remarks.lower():
                        w_type = "POS_WITHDRAWAL"
                    elif "cheque" in s_name_lower or "cheque" in remarks.lower():
                        w_type = "CHEQUE_WITHDRAWAL"
                    elif "aeps" in s_name_lower or "aeps" in remarks.lower():
                        w_type = "AEPS_WITHDRAWAL"
                    elif "cash" in s_name_lower or "cash" in remarks.lower():
                        w_type = "CASH_WITHDRAWAL"

                    # Additional withdrawal location details
                    loc = remarks
                    if "atm id" in " ".join(headers).lower():
                        atm_id_col = self._find_col(headers, ["atm id"])
                        if atm_id_col != -1 and atm_id_col < len(row_cells) and row_cells[atm_id_col]:
                            loc = f"{row_cells[atm_id_col]} | {loc}"

                    all_withdrawals.append({
                        "account_number": src_acc or dst_acc,
                        "withdrawal_type": w_type,
                        "amount": amt,
                        "raw_amount": str(amt),
                        "date": date_val,
                        "raw_date": date_val,
                        "bank_name": bank,
                        "utr_rrn": utr,
                        "location": loc,
                        "remarks": remarks,
                        "sheet_name": sheet_name,
                        "layer_from_sheet": layer,
                        "disputed_amount": disputed,
                    })
                else:
                    # Account-to-account transfer. Only sheets that explicitly
                    # describe transfers (e.g. "Money Transfer to") may create
                    # Money Trail edges, and only with both endpoints present.
                    # Rows from holds / low-value / miscellaneous sheets stay as
                    # evidence records so nothing is silently discarded.
                    if not is_transfer_sheet or not src_acc or not dst_acc:
                        other_records.append({
                            "record_type": sheet_name,
                            "sheet_name": sheet_name,
                            "account_number": src_acc or dst_acc,
                            "utr_rrn": utr,
                            "amount": amt,
                            "disputed_amount": disputed,
                            "date": date_val,
                            "bank_name": bank,
                            "remarks": remarks,
                            "layer_from_doc": layer,
                            "row_text": " | ".join([c for c in row_cells if c]),
                        })
                        continue
                    all_transfers.append({
                        "from_account_number": src_acc,
                        "to_account_number": dst_acc,
                        "bank_name": bank,
                        "raw_layer_from_doc": layer,
                        "ifsc": ifsc,
                        "amount": amt,
                        "raw_amount": str(amt),
                        "disputed_amount": disputed,
                        "raw_date": date_val,
                        "transaction_date": date_val,
                        "utr_rrn": utr,
                        "remarks": remarks,
                        "transaction_type": "ACCOUNT_TRANSFER",
                        "source_sheet": sheet_name,
                    })

        return {
            "transfers": all_transfers,
            "withdrawals": all_withdrawals,
            "other_records": other_records,
            "sections": sections,
            "metadata": {
                "page_count": len(sheets_meta),
                "sheet_names": sheets_meta,
                "document_type": "EXCEL",
                "reported_fraud_amount": reported_fraud_amount,
            },
        }

    def _find_col(self, headers: list[str], keywords: list[str]) -> int:
        for kw in keywords:
            for idx, h in enumerate(headers):
                if kw in h:
                    return idx
        return -1


excel_extractor = ExcelExtractor()
