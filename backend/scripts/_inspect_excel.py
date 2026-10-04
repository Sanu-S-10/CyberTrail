import sys
from pathlib import Path
import openpyxl

xlsx = Path(__file__).resolve().parents[2] / "Reference" / "CR 45-2025 NCRP.xlsx"
wb = openpyxl.load_workbook(xlsx, data_only=True)
for name in wb.sheetnames:
    ws = wb[name]
    print("\n==== SHEET", name, "dims", ws.dimensions)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i > 8:
            break
        vals = [str(c)[:40] if c is not None else "" for c in row[:12]]
        print(i, vals)
